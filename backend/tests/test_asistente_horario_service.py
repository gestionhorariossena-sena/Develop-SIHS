"""Servicio del asistente de programación. Mockea httpx.post para la
parte de IA (mismo patrón que test_ai_classify_document.py) y usa
SQLite en memoria para la parte de catálogo (mismo patrón que
test_horarios_auditoria_cruces.py)."""

import json
import uuid
from datetime import date, time
from io import BytesIO

import openpyxl
import pytest

from app.models.ambiente import Ambiente
from app.models.competencia_formacion import CompetenciaFormacion
from app.models.coordinacion import Coordinacion
from app.models.ficha import Ficha
from app.models.horario import Horario, horario_dia
from app.models.programa import Programa
from app.models.resultado_aprendizaje import ResultadoAprendizaje
from app.models.sede import Sede
from app.models.trimestre import Trimestre
from app.models.usuario import Usuario
from app.services.asistente_horario_service import _elegir_hoja, generar_propuesta, previsualizar_excel


def _crear_tablas_extra(db_session):
    from app.core.database import Base

    Base.metadata.create_all(
        bind=db_session.bind,
        tables=[
            Coordinacion.__table__, Programa.__table__, Trimestre.__table__, Sede.__table__,
            Ambiente.__table__, Ficha.__table__, CompetenciaFormacion.__table__,
            ResultadoAprendizaje.__table__, Horario.__table__, horario_dia,
        ],
    )


def _catalogo_base(db_session, id_ficha=100, codigo_ficha=None):
    # id_ficha (PK interno, SERIAL) y codigo_ficha (el número real de SENA,
    # texto) son columnas DISTINTAS -- default codigo_ficha=str(id_ficha)
    # solo por conveniencia en los tests que no les importa la diferencia;
    # test_previsualizar_excel_reconoce_ficha_existente los desacopla a
    # propósito para no repetir el bug real que encontramos (buscar por PK
    # en vez de por codigoFicha).
    codigo_ficha = codigo_ficha or str(id_ficha)
    db_session.add(Coordinacion(idCoordinacion=1, nombreCoordinacion="Demo"))
    db_session.add(Programa(idPrograma=1, codigoPrograma="P1", nombrePrograma="ADSO", nivelFormacion="Tecnólogo", activo=True, idCoordinacion=1))
    db_session.add(Trimestre(idTrimestre=1, nombre="2026-3", fechaInicio=date(2026, 7, 1), fechaFin=date(2026, 9, 30), estado="activo"))
    db_session.add(Sede(id=1, nombre="Sede Demo", direccion="Calle 1", tipo="principal"))
    db_session.add(Ambiente(id=1, numero_ambiente=101, nombre="Ambiente", tipo_ambiente="regular", estado_ambiente="disponible", sede_id=1))
    db_session.add(Ficha(idFicha=id_ficha, codigoFicha=codigo_ficha, idPrograma=1, idTrimestre=1, idSede=1))
    db_session.add(CompetenciaFormacion(idCompetencia=1, codigo="C1", descripcion="Competencia demo", idPrograma=1))
    db_session.add(ResultadoAprendizaje(idResultado=1, codigo="RA-1", descripcion="Resultado demo", idCompetencia=1, horasAsignadas=20))
    db_session.commit()


def _xlsx_con_encabezado(filas: list[list]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["TÍTULO DEL REPORTE"])
    ws.append([])
    for fila in filas:
        ws.append(fila)
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


class _RespuestaFalsa:
    def __init__(self, cuerpo_json):
        self._cuerpo_json = cuerpo_json

    def json(self):
        return self._cuerpo_json

    def raise_for_status(self):
        pass


def _mock_clasificacion(monkeypatch, mapa: dict[str, tuple]):
    """mapa: {columna_original: (campo, confianza)}"""
    payload = {col: {"campo": campo, "confianza": conf} for col, (campo, conf) in mapa.items()}
    texto_llm = json.dumps(payload)

    def fake_post(url, headers=None, json=None, timeout=None):
        return _RespuestaFalsa({"candidates": [{"content": {"parts": [{"text": texto_llm}]}}]})

    monkeypatch.setattr("app.ai.client.settings.gemini_api_key", "key-de-prueba")
    monkeypatch.setattr("app.ai.client.httpx.post", fake_post)


def _mock_clasificacion_secuencial(monkeypatch, mapas: list[dict[str, tuple]]):
    """Como _mock_clasificacion, pero una respuesta distinta por cada
    llamada en orden -- para cuando se clasifican dos hojas (archivo
    principal + complementario), cada una con sus propias columnas."""
    textos = [
        json.dumps({col: {"campo": campo, "confianza": conf} for col, (campo, conf) in mapa.items()})
        for mapa in mapas
    ]
    llamadas = iter(textos)

    def fake_post(url, headers=None, json=None, timeout=None):
        texto = next(llamadas)
        return _RespuestaFalsa({"candidates": [{"content": {"parts": [{"text": texto}]}}]})

    monkeypatch.setattr("app.ai.client.settings.gemini_api_key", "key-de-prueba")
    monkeypatch.setattr("app.ai.client.httpx.post", fake_post)


def test_previsualizar_excel_marca_ficha_inexistente(db_session, monkeypatch):
    _crear_tablas_extra(db_session)
    contenido = _xlsx_con_encabezado([["FICHA", "PROGRAMA"], [999999, "ADSO"]])
    _mock_clasificacion(monkeypatch, {"FICHA": ("ficha", 1.0), "PROGRAMA": ("programa", 1.0)})

    resultado = previsualizar_excel(db_session, contenido, "archivo.xlsx")

    assert resultado.totalFilas == 1
    assert resultado.filas[0].codigoFicha == "999999"
    assert resultado.filas[0].idFicha is None
    assert resultado.filas[0].fichaExiste is False
    assert "no existe" in resultado.filas[0].advertencia


def test_previsualizar_excel_sin_columna_de_ficha_da_advertencia_general(db_session, monkeypatch):
    # Caso real encontrado en pruebas: un formato de matriz/pivote donde
    # ninguna columna se clasifica como "ficha" -- en vez de 40 filas con
    # el mismo mensaje sin contexto, un solo aviso a nivel de archivo.
    _crear_tablas_extra(db_session)
    contenido = _xlsx_con_encabezado([["ITEMS", "TEMAS_"], ["a", "TEMAS_7_TRM_2996161_(DM)_ALGO"]])
    _mock_clasificacion(monkeypatch, {"ITEMS": (None, 0.0), "TEMAS_": (None, 0.0)})

    resultado = previsualizar_excel(db_session, contenido, "archivo.xlsx")

    assert resultado.advertenciaGeneral is not None
    assert "formato de matriz o pivote" in resultado.advertenciaGeneral


def test_previsualizar_excel_reconoce_ficha_existente(db_session, monkeypatch):
    # idFicha (PK interno) != codigoFicha (número real de SENA) a propósito
    # -- si el servicio buscara por PK en vez de por codigoFicha, esta
    # ficha real "2895566" no se encontraría aunque exista en la BD.
    _crear_tablas_extra(db_session)
    _catalogo_base(db_session, id_ficha=7, codigo_ficha="2895566")
    contenido = _xlsx_con_encabezado([["FICHA", "PROGRAMA"], [2895566, "ADSO"]])
    _mock_clasificacion(monkeypatch, {"FICHA": ("ficha", 1.0), "PROGRAMA": ("programa", 1.0)})

    resultado = previsualizar_excel(db_session, contenido, "archivo.xlsx")

    assert resultado.filas[0].codigoFicha == "2895566"
    assert resultado.filas[0].fichaExiste is True
    assert resultado.filas[0].idFicha == 7
    assert resultado.filas[0].advertencia is None


def test_generar_propuesta_sin_resultados_pendientes(db_session):
    _crear_tablas_extra(db_session)
    _catalogo_base(db_session, id_ficha=100)
    # Ya existe un horario activo para el único resultado de esta ficha.

    instructor_id = uuid.uuid4()
    db_session.add(Usuario(idUsuario=instructor_id, nombre="Ana", email="ana@demo.sihs", tipoContrato="contrato", estado="activo"))
    db_session.add(Horario(
        idHorario=1, horaInicio=time(7, 0), horaFin=time(9, 0), idJornada=1, idTrimestre=1,
        idAmbiente=1, idInstructor=instructor_id, idFicha=100, idResultado=1, activo=True,
    ))
    db_session.commit()

    resultado = generar_propuesta(db_session, id_trimestre=1, ids_ficha=[100], jornada="MAÑANA")

    assert resultado.bloques == []
    assert resultado.factible is True
    assert "ya tienen todos sus resultados" in resultado.mensaje


def test_generar_propuesta_programa_sin_resultados_definidos_da_mensaje_distinto(db_session):
    # Caso real encontrado en pruebas: crear una ficha con un Programa
    # nuevo (sin ninguna CompetenciaFormacion/ResultadoAprendizaje
    # cargada todavía) no debe confundirse con "ya está todo programado"
    # -- son causas y remedios distintos.
    _crear_tablas_extra(db_session)
    db_session.add(Coordinacion(idCoordinacion=1, nombreCoordinacion="Demo"))
    db_session.add(Programa(idPrograma=1, codigoPrograma="P1", nombrePrograma="Programa nuevo", activo=True, idCoordinacion=1))
    db_session.add(Trimestre(idTrimestre=1, nombre="2026-3", fechaInicio=date(2026, 7, 1), fechaFin=date(2026, 9, 30), estado="activo"))
    db_session.add(Sede(id=1, nombre="Sede Demo", direccion="Calle 1", tipo="principal"))
    db_session.add(Ambiente(id=1, numero_ambiente=101, nombre="Ambiente", tipo_ambiente="regular", estado_ambiente="disponible", sede_id=1))
    db_session.add(Ficha(idFicha=200, codigoFicha="200", idPrograma=1, idTrimestre=1, idSede=1))
    instructor_id = uuid.uuid4()
    db_session.add(Usuario(idUsuario=instructor_id, nombre="Ana", email="ana2@demo.sihs", tipoContrato="contrato", estado="activo"))
    db_session.commit()

    resultado = generar_propuesta(db_session, id_trimestre=1, ids_ficha=[200], jornada="MAÑANA")

    assert resultado.bloques == []
    assert resultado.factible is True
    assert "no tiene resultados de aprendizaje" in resultado.mensaje


def test_generar_propuesta_genera_un_bloque_real(db_session):
    _crear_tablas_extra(db_session)
    _catalogo_base(db_session, id_ficha=100)

    instructor_id = uuid.uuid4()
    db_session.add(Usuario(idUsuario=instructor_id, nombre="Ana", email="ana@demo.sihs", tipoContrato="contrato", estado="activo"))
    db_session.commit()

    resultado = generar_propuesta(db_session, id_trimestre=1, ids_ficha=[100], jornada="MAÑANA")

    assert resultado.factible is True
    assert len(resultado.bloques) == 1
    bloque = resultado.bloques[0]
    assert bloque.idFicha == 100
    assert bloque.idResultado == 1
    assert bloque.instructorNombre == "Ana"
    assert bloque.idAmbiente == 1


def test_generar_propuesta_filtra_por_fase_actual_de_la_ficha(db_session):
    # Caso real que colgaba el asistente: un programa con currículo de
    # varias fases (TRIM I..IV) y una ficha que solo debería programar la
    # fase en la que va, no las demás. resultado 1 = fase 1 (coincide con
    # faseActual de la ficha), resultado 2 = fase 2 (no debe aparecer).
    _crear_tablas_extra(db_session)
    _catalogo_base(db_session, id_ficha=100)
    db_session.query(Ficha).filter(Ficha.idFicha == 100).update({"faseActual": 1})
    db_session.add(ResultadoAprendizaje(idResultado=2, codigo="RA-2", descripcion="Resultado fase 2", idCompetencia=1, numeroFase=2))
    db_session.query(ResultadoAprendizaje).filter(ResultadoAprendizaje.idResultado == 1).update({"numeroFase": 1})
    instructor_id = uuid.uuid4()
    db_session.add(Usuario(idUsuario=instructor_id, nombre="Ana", email="ana@demo.sihs", tipoContrato="contrato", estado="activo"))
    db_session.commit()

    resultado = generar_propuesta(db_session, id_trimestre=1, ids_ficha=[100], jornada="MAÑANA")

    assert resultado.factible is True
    assert len(resultado.bloques) == 1
    assert resultado.bloques[0].idResultado == 1


def test_generar_propuesta_infactible_da_mensaje_con_diagnostico(db_session):
    # Volumen real que hacía fallar el solver: varios resultados
    # pendientes compitiendo por muy pocos instructores/ambientes. El
    # mensaje debe explicar la capacidad, no solo decir "no se encontró".
    _crear_tablas_extra(db_session)
    db_session.add(Coordinacion(idCoordinacion=1, nombreCoordinacion="Demo"))
    db_session.add(Programa(idPrograma=1, codigoPrograma="P1", nombrePrograma="ADSO", activo=True, idCoordinacion=1))
    db_session.add(Trimestre(idTrimestre=1, nombre="2026-3", fechaInicio=date(2026, 7, 1), fechaFin=date(2026, 9, 30), estado="activo"))
    db_session.add(Sede(id=1, nombre="Sede Demo", direccion="Calle 1", tipo="principal"))
    db_session.add(Ambiente(id=1, numero_ambiente=101, nombre="Ambiente", tipo_ambiente="regular", estado_ambiente="disponible", sede_id=1))
    db_session.add(Ficha(idFicha=100, codigoFicha="100", idPrograma=1, idTrimestre=1, idSede=1))
    db_session.add(CompetenciaFormacion(idCompetencia=1, codigo="C1", descripcion="Competencia demo", idPrograma=1))
    # Un solo instructor, un solo ambiente: como mucho caben 5 días × 3
    # franjas = 15 bloques en la semana. 20 resultados pendientes no caben.
    for i in range(1, 21):
        db_session.add(ResultadoAprendizaje(idResultado=i, codigo=f"RA-{i}", descripcion=f"Resultado {i}", idCompetencia=1))
    instructor_id = uuid.uuid4()
    db_session.add(Usuario(idUsuario=instructor_id, nombre="Ana", email="ana@demo.sihs", tipoContrato="contrato", estado="activo"))
    db_session.commit()

    resultado = generar_propuesta(db_session, id_trimestre=1, ids_ficha=[100], jornada="MAÑANA")

    assert resultado.factible is False
    assert "20 resultado" in resultado.mensaje
    assert "capacidad" in resultado.mensaje
    assert "Reduce cuántas fichas" in resultado.mensaje


def test_generar_propuesta_jornada_invalida_lanza_value_error(db_session):
    _crear_tablas_extra(db_session)
    _catalogo_base(db_session, id_ficha=100)

    with pytest.raises(ValueError, match="no reconocida"):
        generar_propuesta(db_session, id_trimestre=1, ids_ficha=[100], jornada="MADRUGADA")


def test_elegir_hoja_prefiere_la_hoja_con_mas_fichas_numericas():
    # Caso real: PLANEACION tiene una columna que se LLAMA "ficha" pero
    # con valores compuestos (texto), FICHAS tiene el número limpio.
    wb = openpyxl.Workbook()
    hoja_mala = wb.active
    hoja_mala.title = "PLANEACION"
    hoja_mala.append(["ficha", "otro"])
    hoja_mala.append(["TEMAS_7_TRM_2996161_(DM)_ALGO", "x"])

    hoja_buena = wb.create_sheet("FICHAS")
    hoja_buena.append(["FICHA", "NIVEL"])
    hoja_buena.append([2996161, "TECNÓLOGO"])
    hoja_buena.append([2996202, "TECNÓLOGO"])

    ws, _, columna = _elegir_hoja(wb)

    assert ws.title == "FICHAS"
    assert columna == "FICHA"


def test_elegir_hoja_compara_dos_columnas_ficha_en_la_misma_hoja():
    # El bug real: PE-04 trae IDENTIFICADOR_FICHA e
    # IDENTIFICADOR_UNICO_FICHA en la MISMA hoja -- hay que comparar entre
    # ellas, no quedarse con la primera que aparezca.
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["FICHA_COMPUESTA", "FICHA_REAL"])
    ws.append(["ABC-100-XYZ", 100])
    ws.append(["ABC-200-XYZ", 200])

    _, _, columna = _elegir_hoja(wb)

    assert columna == "FICHA_REAL"


def test_previsualizar_excel_corrige_columna_ficha_mal_clasificada_por_ia(db_session, monkeypatch):
    _crear_tablas_extra(db_session)
    _catalogo_base(db_session, id_ficha=7, codigo_ficha="100")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["FICHA_COMPUESTA", "FICHA_REAL", "PROGRAMA"])
    ws.append(["ABC-100-XYZ", 100, "ADSO"])
    buffer = BytesIO()
    wb.save(buffer)
    contenido = buffer.getvalue()

    # La IA (mockeada) "se equivoca" y mapea "ficha" a la columna
    # compuesta -- el override mecánico de _elegir_hoja debe corregirlo.
    _mock_clasificacion(
        monkeypatch,
        {"FICHA_COMPUESTA": ("ficha", 0.9), "FICHA_REAL": (None, 0.0), "PROGRAMA": ("programa", 1.0)},
    )

    resultado = previsualizar_excel(db_session, contenido, "archivo.xlsx")

    assert resultado.filas[0].codigoFicha == "100"
    assert resultado.filas[0].fichaExiste is True


def test_previsualizar_excel_cruza_con_archivo_complementario(db_session, monkeypatch):
    _crear_tablas_extra(db_session)
    principal = _xlsx_con_encabezado([["FICHA", "PROGRAMA"], [100, "ADSO"]])
    complementario = _xlsx_con_encabezado(
        [["FICHA", "NIVEL", "COORDINACION", "CODIGO"], [100, "Tecnólogo", "Teleinformática", "228106"]]
    )

    _mock_clasificacion_secuencial(
        monkeypatch,
        [
            {"FICHA": ("ficha", 1.0), "PROGRAMA": ("programa", 1.0)},
            {
                "FICHA": ("ficha", 1.0),
                "NIVEL": ("nivel_formacion", 0.95),
                "COORDINACION": ("coordinacion", 0.9),
                "CODIGO": ("codigo_programa", 0.9),
            },
        ],
    )

    resultado = previsualizar_excel(db_session, principal, "principal.xlsx", complementario, "complementario.xlsx")

    assert resultado.archivoComplementario == "complementario.xlsx"
    fila = resultado.filas[0]
    assert fila.nivelFormacion == "Tecnólogo"
    assert fila.coordinacion == "Teleinformática"
    assert fila.codigoPrograma == "228106"
