"""Importa contenido curricular (competencias + resultados de
aprendizaje) real desde el "Formato Planeación Pedagógica" de SENA --
ver PLAN_INTEGRACION_IA.md. Sin IA a propósito: el formato tiene
encabezados fijos ("COMPETENCIA" / "RESULTADOS DE APRENDIZAJE"), no hace
falta clasificación semántica para reconocerlo -- validado contra dos
archivos reales del usuario (27 resultados / 7 competencias, y una
variante similar).

Solo `previsualizar_curriculo` -- no persiste nada. El frontend confirma
llamando a POST /competencias-formacion/ y POST /resultados-aprendizaje/
(ya existentes) por cada fila, igual que ImportarArchivo.tsx hace para
otros catálogos -- no hace falta un endpoint de "crear todo junto"
nuevo.
"""

from io import BytesIO

import openpyxl

from app.schemas.curriculo import CompetenciaExtraida, PreviewCurriculoResponse, ResultadoExtraido

_FILAS_A_BUSCAR_ENCABEZADO = 30


def _fila_encabezado_curriculo(ws) -> tuple[int, int, int, int | None] | None:
    """Busca la fila con los encabezados "COMPETENCIA" y "RESULTADOS DE
    APRENDIZAJE" (pueden estar en cualquier columna). Devuelve
    (fila, columna_competencia, columna_resultados, columna_horas) o
    None si no aparecen -- este archivo no es de este formato."""
    for fila in range(1, min(ws.max_row, _FILAS_A_BUSCAR_ENCABEZADO) + 1):
        valores = [str(c.value).strip().upper() if c.value not in (None, "") else "" for c in ws[fila]]
        if "COMPETENCIA" not in valores:
            continue
        idx_resultados = next((i for i, v in enumerate(valores) if "RESULTADOS DE APRENDIZAJE" in v), None)
        if idx_resultados is None:
            continue
        idx_competencia = valores.index("COMPETENCIA")
        idx_horas = next((i for i, v in enumerate(valores) if "DURACIÓN" in v or "DURACION" in v), None)
        return fila, idx_competencia, idx_resultados, idx_horas
    return None


def previsualizar_curriculo(contenido: bytes, nombre_archivo: str) -> PreviewCurriculoResponse:
    wb = openpyxl.load_workbook(BytesIO(contenido), data_only=True)
    ws = wb.worksheets[0]

    encabezado = _fila_encabezado_curriculo(ws)
    if encabezado is None:
        raise ValueError(
            'No se encontraron las columnas "COMPETENCIA" y "RESULTADOS DE APRENDIZAJE" en este '
            "archivo -- no parece ser un Formato de Planeación Pedagógica."
        )
    fila_encabezado, idx_competencia, idx_resultados, idx_horas = encabezado

    # COMPETENCIA suele venir en celdas combinadas -- el valor solo
    # aparece en la primera fila del grupo, las siguientes están vacías
    # (mismo patrón que Excel real de SENA en ambos archivos probados).
    resultados_por_competencia: dict[str, list[ResultadoExtraido]] = {}
    orden_competencias: list[str] = []
    competencia_actual: str | None = None

    for fila_valores in ws.iter_rows(min_row=fila_encabezado + 1, values_only=True):
        valor_competencia = fila_valores[idx_competencia] if idx_competencia < len(fila_valores) else None
        valor_resultado = fila_valores[idx_resultados] if idx_resultados < len(fila_valores) else None

        if valor_competencia not in (None, ""):
            competencia_actual = str(valor_competencia).strip()
            if competencia_actual not in resultados_por_competencia:
                resultados_por_competencia[competencia_actual] = []
                orden_competencias.append(competencia_actual)

        if valor_resultado in (None, "") or competencia_actual is None:
            continue

        horas = None
        if idx_horas is not None and idx_horas < len(fila_valores):
            valor_horas = fila_valores[idx_horas]
            if isinstance(valor_horas, (int, float)):
                horas = int(valor_horas)

        resultados_por_competencia[competencia_actual].append(
            ResultadoExtraido(descripcion=str(valor_resultado).strip(), horasAsignadas=horas)
        )

    competencias = [
        CompetenciaExtraida(descripcion=nombre, resultados=resultados_por_competencia[nombre])
        for nombre in orden_competencias
        if resultados_por_competencia[nombre]
    ]

    return PreviewCurriculoResponse(
        nombreArchivo=nombre_archivo,
        hoja=ws.title,
        competencias=competencias,
        totalCompetencias=len(competencias),
        totalResultados=sum(len(c.resultados) for c in competencias),
    )
