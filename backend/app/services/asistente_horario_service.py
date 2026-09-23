"""Servicio del asistente de programación (wizard de 4 pasos). Dos
responsabilidades separadas a propósito:

1. `previsualizar_excel`: solo lee y clasifica un Excel real, nunca
   escribe nada -- pasos 1-2 del wizard ("Subir archivo" / "Así lo
   entendimos").
2. `generar_propuesta`: solo trabaja con fichas que YA existen en el
   catálogo de SIHS -- si el Excel trae una ficha nueva, el import la
   señala como pendiente en vez de inventarla en la BD. Programar una
   ficha que no existe todavía en el sistema no es un caso de este
   asistente; primero se crea en Fichas, después se programa acá. Paso 3
   ("Así quedaría el horario").

Ninguna de las dos persiste horarios -- eso lo hace
`HorarioService.crear`, llamado directo desde el paso 4 (confirmar) del
frontend, reusando la validación de cruces que ya existe.
"""

import re
import time
from datetime import date, datetime
from io import BytesIO

import openpyxl
from sqlalchemy.orm import Session

from app.ai.tasks.classify_document import clasificar_columnas
from app.models.ambiente import Ambiente
from app.models.competencia_formacion import CompetenciaFormacion
from app.models.ficha import Ficha
from app.models.horario import Horario, horario_dia
from app.models.resultado_aprendizaje import ResultadoAprendizaje
from app.models.usuario import Usuario
from app.schemas.asistente_horario import (
    BloquePropuesto,
    ColumnaClasificada,
    FilaImportada,
    GenerarPropuestaResponse,
    ImportarExcelPreviewResponse,
)
from app.scheduling.generator import FRANJAS_POR_JORNADA, NecesidadHorario, generar_horario

CONFIANZA_MINIMA = 0.7
MAX_FILAS_PREVIA = 50

# Ver app/scheduling/generator.py -- mismo catálogo de 3 jornadas.
JORNADAS_VALIDAS = ("MAÑANA", "TARDE", "NOCHE")
_ID_JORNADA_POR_CLAVE = {"MAÑANA": 1, "TARDE": 2, "NOCHE": 3}


def _fila_con_mas_datos(ws, filas_a_revisar: int = 10) -> int:
    """Heurística simple: la fila de encabezado real suele tener más
    celdas no vacías que las filas de título que suelen venir antes (ver
    LIDERES DE FICHA 2026_pruebas.xlsx: título en la fila 1, encabezado
    real en la fila 3)."""
    mejor_fila, mejor_conteo = 1, 0
    for fila in range(1, min(ws.max_row, filas_a_revisar) + 1):
        conteo = sum(1 for celda in ws[fila] if celda.value not in (None, "", "\xa0"))
        if conteo > mejor_conteo:
            mejor_fila, mejor_conteo = fila, conteo
    return mejor_fila


def _elegir_hoja(wb, codigos_objetivo: set[str] | None = None) -> tuple:
    """Si el archivo tiene varias hojas, no asume que la útil es la
    primera -- ver PROGRAMACIÓN CGMLTI I TRM 2026 (4).xlsx, donde la hoja
    limpia de fichas (FICHAS) es la #14 de 15, y hojas anteriores como
    PLANEACION también tienen una columna que se LLAMA "ficha" pero con
    valores compuestos, no el número solo.

    Heurística barata (sin IA, solo para elegir la hoja): busca en cada
    hoja TODOS los encabezados que contengan "ficha" (puede haber más de
    uno en la misma hoja -- ver PE-04 de SOFIA Plus, que trae a la vez
    "IDENTIFICADOR_FICHA" y "IDENTIFICADOR_UNICO_FICHA" con un prefijo
    extra, y semánticamente ambos "suenan" a ficha) y cuenta, para cada
    uno, cuántas filas de datos son útiles. También se devuelve cuál era
    esa columna ganadora: es una señal más confiable que la IA para el
    campo "ficha" específicamente (ver _clasificar_hoja).

    Con `codigos_objetivo` (los codigoFicha ya conocidos, del archivo
    principal): puntúa por CUÁNTOS de esos códigos reales aparecen en
    cada columna candidata -- la pregunta que de verdad importa para el
    archivo complementario es "¿esta hoja trae datos de las fichas que
    ya tengo?", no "¿qué tan numérica se ve esta columna?". Sin
    `codigos_objetivo` (archivo principal, donde codigoFicha es lo que se
    está extrayendo, no algo ya conocido) usa la pureza numérica como
    proxy barato.

    Encontrado en vivo el 2026-09-13: con la pureza numérica sola, PE-04
    (un volcado crudo de SOFIA Plus con miles de IDs puramente numéricos,
    pero ninguno coincidiendo con el número real de ficha) le ganaba a
    FICHAS (la hoja correcta, con menos filas y algunas con sufijo de
    letra tipo "3228970 A" que rompía el isdigit()) -- 0 de 50 fichas
    del lote recibían faseActual del complementario porque se estaba
    leyendo la hoja equivocada, no por ningún problema de datos."""
    mejor_ws, mejor_fila_encabezado, mejor_puntaje, mejor_columna_ficha = (
        wb.worksheets[0], _fila_con_mas_datos(wb.worksheets[0]), -1, None,
    )
    for ws in wb.worksheets:
        fila_encabezado = _fila_con_mas_datos(ws)
        encabezados = [str(c.value).strip() if c.value not in (None, "") else "" for c in ws[fila_encabezado]]
        indices_candidatos = [i for i, h in enumerate(encabezados) if "ficha" in h.lower()]
        if not indices_candidatos:
            continue

        puntajes = dict.fromkeys(indices_candidatos, 0)
        # Con un objetivo concreto que buscar, vale la pena recorrer toda
        # la hoja (no solo una muestra) -- el archivo complementario puede
        # tener miles de filas y el código que buscamos puede estar en
        # cualquier parte, no necesariamente entre las primeras 30.
        max_row = None if codigos_objetivo else fila_encabezado + 30
        for fila_valores in ws.iter_rows(min_row=fila_encabezado + 1, max_row=max_row, values_only=True):
            for idx_ficha in indices_candidatos:
                if idx_ficha < len(fila_valores):
                    valor = fila_valores[idx_ficha]
                    if valor is None:
                        continue
                    if codigos_objetivo:
                        if _codigo_ficha_desde_texto(str(valor)) in codigos_objetivo:
                            puntajes[idx_ficha] += 1
                    elif str(valor).strip().isdigit():
                        puntajes[idx_ficha] += 1

        idx_mejor_de_la_hoja = max(puntajes, key=lambda i: puntajes[i])
        puntaje = puntajes[idx_mejor_de_la_hoja]
        if puntaje > mejor_puntaje:
            mejor_ws, mejor_fila_encabezado, mejor_puntaje = ws, fila_encabezado, puntaje
            mejor_columna_ficha = encabezados[idx_mejor_de_la_hoja]
    return mejor_ws, mejor_fila_encabezado, mejor_columna_ficha


def _clasificar_hoja(wb, codigos_objetivo: set[str] | None = None) -> tuple:
    """Elige la mejor hoja (_elegir_hoja) y clasifica sus encabezados con
    IA una sola vez -- devuelve todo lo necesario para leer filas de ella
    sin volver a llamar a la IA por cada uso. `codigos_objetivo` se pasa
    tal cual a _elegir_hoja -- ver su docstring."""
    ws, fila_encabezado, columna_ficha_detectada = _elegir_hoja(wb, codigos_objetivo)
    encabezados_crudos = [c.value for c in ws[fila_encabezado]]
    idx = {str(nombre).strip(): i for i, nombre in enumerate(encabezados_crudos) if nombre not in (None, "", "\xa0")}
    encabezados = list(idx.keys())

    clasificacion = clasificar_columnas(encabezados) if encabezados else None
    columnas: list[ColumnaClasificada] = []
    campo_a_columna: dict[str, str] = {}
    if clasificacion:
        for original, c in clasificacion.root.items():
            columnas.append(ColumnaClasificada(columnaOriginal=original, campo=c.campo, confianza=c.confianza))
            if c.campo and c.confianza >= CONFIANZA_MINIMA:
                campo_a_columna[c.campo] = original

    # Para "ficha" específicamente, la columna que _elegir_hoja ya midió
    # por densidad de valores numéricos reales gana sobre lo que haya
    # dicho la IA por semántica del nombre -- ver docstring de _elegir_hoja.
    if columna_ficha_detectada and columna_ficha_detectada in idx:
        campo_a_columna["ficha"] = columna_ficha_detectada

    # "TRI"/"TRM" son casi siempre fase_actual del pénsum en los archivos
    # reales que ha traído el coordinador (ver CAMPOS_CONOCIDOS_SIHS y
    # PLAN_INTEGRACION_IA.md) -- pero la IA los confunde con "trimestre"
    # (el periodo académico, un campo real distinto) porque el nombre
    # corto se le parece. Encontrado en vivo el 2026-09-13 con la hoja
    # FICHAS real: Gemini clasificó "TRI" como trimestre con 0.9 de
    # confianza, y faseActual quedaba vacío para TODAS las fichas del
    # complementario sin que hubiera ningún problema en los datos. Se
    # sobreescribe la clasificación de la IA para este encabezado exacto
    # en vez de esperar a que adivine bien la próxima vez.
    for encabezado in encabezados:
        if encabezado.strip().upper() in ("TRI", "TRM"):
            campo_a_columna["fase_actual"] = encabezado
            break

    return ws, fila_encabezado, idx, campo_a_columna, columnas


def _valor_texto(fila_valores: tuple, idx: dict, campo_a_columna: dict, campo: str) -> str | None:
    columna = campo_a_columna.get(campo)
    if not columna or columna not in idx or idx[columna] >= len(fila_valores):
        return None
    valor = fila_valores[idx[columna]]
    return str(valor).strip() if valor not in (None, "", "\xa0") else None


_PATRON_FICHA_CON_LETRA = re.compile(r"^(\d+)\s*([A-Za-z])$")
_PATRON_FICHA_ENTRE_PARENTESIS = re.compile(r"\((\d+)\)")


def _codigo_simple_con_o_sin_letra(segmento: str) -> str | None:
    """Un segmento (ya sin guiones) que es o un número puro, o un número
    con letra distintiva -- ver _codigo_ficha_desde_texto."""
    segmento = segmento.strip()
    con_letra = _PATRON_FICHA_CON_LETRA.match(segmento)
    if con_letra:
        return f"{con_letra.group(1)}{con_letra.group(2).upper()}"
    return segmento if segmento.isdigit() else None


def _codigo_ficha_desde_texto(texto: str) -> str | None:
    """El código real de SENA, como texto.

    Dos formatos reales que se pueden combinar, y no hay que confundir:
    - Letra distintiva (ej. "3228973A" / "3228973B", o con espacio
      "3171242 A"/"3171242 B"): son DOS FICHAS DIFERENTES que comparten
      número base -- se conserva la letra (normalizada, sin el espacio)
      como parte del codigoFicha real. Confirmado con el usuario.
    - Guion de unificación (ej. "3171645-65-668", unificada con
      3171665, o "3171667-668"): es la MISMA ficha, física y
      administrativamente unida con otra -- se usa el segmento de la
      IZQUIERDA (antes del primer guion) como el código real.

    Ambos combinados: "3228970 A - B" es la ficha "3228970 A" unificada
    con su par "B" -- el segmento de la izquierda ("3228970 A") todavía
    tiene la letra distintiva, así que hay que re-aplicar el patrón de
    letra sobre ÉL, no solo comprobar que sea puramente numérico.

    Un tercer formato, distinto de los dos anteriores: número entre
    paréntesis (ej. "3311985 (3288277)") -- confirmado con el usuario:
    se usa el que está DENTRO del paréntesis como el código real (al
    revés que el guion, donde se usa el de la izquierda).

    Devuelve None si nada de esto calza (dato realmente irreconocible,
    ej. texto libre)."""
    texto = texto.strip()
    entre_parentesis = _PATRON_FICHA_ENTRE_PARENTESIS.search(texto)
    if entre_parentesis:
        return entre_parentesis.group(1)
    directo = _codigo_simple_con_o_sin_letra(texto)
    if directo:
        return directo
    return _codigo_simple_con_o_sin_letra(texto.split("-")[0])


# Solo para "fase_actual": el Excel real la trae de dos formas -- número
# entero plano (columna "TRI" de PROGRAMACIÓN CGMLTI) o número romano
# (columna "TRM" de LIDERES DE FICHA, valores reales vistos: I..VII).
_ROMANOS_A_NUMERO = {"I": 1, "II": 2, "III": 3, "IV": 4, "V": 5, "VI": 6, "VII": 7}


def _valor_entero(fila_valores: tuple, idx: dict, campo_a_columna: dict, campo: str) -> int | None:
    columna = campo_a_columna.get(campo)
    if not columna or columna not in idx or idx[columna] >= len(fila_valores):
        return None
    valor = fila_valores[idx[columna]]
    if isinstance(valor, bool):
        return None
    if isinstance(valor, (int, float)):
        return int(valor)
    if isinstance(valor, str):
        texto = valor.strip()
        if texto.isdigit():
            return int(texto)
        if campo == "fase_actual" and texto.upper() in _ROMANOS_A_NUMERO:
            return _ROMANOS_A_NUMERO[texto.upper()]
    return None


_FORMATOS_FECHA_TEXTO = ("%d/%m/%Y", "%Y-%m-%d")


def _valor_fecha(fila_valores: tuple, idx: dict, campo_a_columna: dict, campo: str) -> date | None:
    columna = campo_a_columna.get(campo)
    if not columna or columna not in idx or idx[columna] >= len(fila_valores):
        return None
    valor = fila_valores[idx[columna]]
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    if isinstance(valor, str):
        # Algunos exports oficiales (ej. PE-04 de SOFIA Plus) guardan la
        # fecha como texto "DD/MM/YYYY", no como celda de fecha real.
        for formato in _FORMATOS_FECHA_TEXTO:
            try:
                return datetime.strptime(valor.strip(), formato).date()
            except ValueError:
                continue
    return None


def _datos_complementarios_por_ficha(
    contenido: bytes, codigos_objetivo: set[str] | None = None
) -> tuple[dict[str, dict], str]:
    """Lee un archivo complementario y devuelve {codigoFicha: {...}} con
    lo que traiga de nivel/coordinación/fechas -- para cruzar con el
    archivo principal por número de ficha. LIDERES DE FICHA no trae estos
    campos, pero por ejemplo la hoja FICHAS de PROGRAMACIÓN CGMLTI sí.

    `codigos_objetivo` (los codigoFicha ya leídos del archivo principal)
    se pasa a _elegir_hoja para que elija la hoja que de verdad trae esos
    códigos, en vez de la que tenga la columna más "numéricamente pura"
    -- ver su docstring para el caso real que esto arregla."""
    wb = openpyxl.load_workbook(BytesIO(contenido), data_only=True)
    ws, fila_encabezado, idx, campo_a_columna, _ = _clasificar_hoja(wb, codigos_objetivo)

    datos: dict[str, dict] = {}
    for fila_valores in ws.iter_rows(min_row=fila_encabezado + 1, values_only=True):
        ficha_texto = _valor_texto(fila_valores, idx, campo_a_columna, "ficha")
        codigo_ficha_extra = _codigo_ficha_desde_texto(ficha_texto) if ficha_texto else None
        if not codigo_ficha_extra:
            continue
        datos[codigo_ficha_extra] = {
            "nivelFormacion": _valor_texto(fila_valores, idx, campo_a_columna, "nivel_formacion"),
            "coordinacion": _valor_texto(fila_valores, idx, campo_a_columna, "coordinacion"),
            "codigoPrograma": _valor_texto(fila_valores, idx, campo_a_columna, "codigo_programa"),
            "fechaInicioLectiva": _valor_fecha(fila_valores, idx, campo_a_columna, "fecha_inicio_lectiva"),
            "fechaFinLectiva": _valor_fecha(fila_valores, idx, campo_a_columna, "fecha_fin_lectiva"),
            "fechaFinProductiva": _valor_fecha(fila_valores, idx, campo_a_columna, "fecha_fin_productiva"),
            "faseActual": _valor_entero(fila_valores, idx, campo_a_columna, "fase_actual"),
        }
    return datos, ws.title


def previsualizar_excel(
    db: Session,
    contenido: bytes,
    nombre_archivo: str,
    contenido_complementario: bytes | None = None,
    nombre_complementario: str | None = None,
) -> ImportarExcelPreviewResponse:
    wb = openpyxl.load_workbook(BytesIO(contenido), data_only=True)
    ws, fila_encabezado, idx, campo_a_columna, columnas = _clasificar_hoja(wb)

    datos_complementarios: dict[str, dict] = {}
    hoja_complementaria: str | None = None
    if contenido_complementario:
        # Los codigoFicha que ya trae el archivo principal -- se usan
        # para que _elegir_hoja escoja la hoja del complementario que de
        # verdad tiene datos de ESTAS fichas (ver su docstring), no la
        # que tenga la columna más "numéricamente pura" a ciegas.
        codigos_ficha_principal: set[str] = set()
        for fila_valores in ws.iter_rows(min_row=fila_encabezado + 1, values_only=True):
            ficha_texto = _valor_texto(fila_valores, idx, campo_a_columna, "ficha")
            codigo = _codigo_ficha_desde_texto(ficha_texto) if ficha_texto else None
            if codigo:
                codigos_ficha_principal.add(codigo)
        datos_complementarios, hoja_complementaria = _datos_complementarios_por_ficha(
            contenido_complementario, codigos_ficha_principal
        )

    filas: list[FilaImportada] = []
    for numero_fila, fila_valores in enumerate(
        ws.iter_rows(min_row=fila_encabezado + 1, values_only=True), start=fila_encabezado + 1
    ):
        if all(v in (None, "", "\xa0") for v in fila_valores):
            continue

        ficha_texto = _valor_texto(fila_valores, idx, campo_a_columna, "ficha")
        codigo_ficha: str | None = None
        id_ficha: int | None = None
        fase_actual_en_bd: int | None = None
        advertencia: str | None = None
        ficha_existe = False

        if ficha_texto is None:
            advertencia = "No se reconoció la columna de ficha en esta fila."
        else:
            codigo_ficha = _codigo_ficha_desde_texto(ficha_texto)
            if codigo_ficha is None:
                # Dato real irreconocible (texto libre, no un número ni
                # siquiera en el primer segmento antes de un guion).
                advertencia = f"Ficha en formato no reconocido ({ficha_texto!r}) -- requiere revisión manual."
            else:
                # codigoFicha es texto (el número real de SENA), NO el
                # idFicha interno -- son columnas distintas, nunca hay que
                # buscar por PK acá.
                ficha_db = db.query(Ficha).filter(Ficha.codigoFicha == codigo_ficha).first()
                if ficha_db:
                    ficha_existe = True
                    id_ficha = ficha_db.idFicha
                    fase_actual_en_bd = ficha_db.faseActual
                else:
                    advertencia = f"La ficha {codigo_ficha} no existe todavía en el catálogo de SIHS."

        extra = datos_complementarios.get(codigo_ficha, {}) if codigo_ficha else {}

        # Estos campos pueden venir directo en el archivo principal (ej.
        # "NIVEL" en LIDERES DE FICHA) -- antes solo se leían del
        # complementario, así que un archivo único que sí trae "nivel"
        # igual pedía crear el programa a mano por falta de ese dato. El
        # complementario queda como respaldo cuando el principal no lo trae.
        filas.append(
            FilaImportada(
                fila=numero_fila,
                codigoFicha=codigo_ficha,
                fichaExiste=ficha_existe,
                idFicha=id_ficha,
                programa=_valor_texto(fila_valores, idx, campo_a_columna, "programa"),
                jornada=_valor_texto(fila_valores, idx, campo_a_columna, "jornada"),
                instructorNombre=_valor_texto(fila_valores, idx, campo_a_columna, "instructor"),
                advertencia=advertencia,
                nivelFormacion=_valor_texto(fila_valores, idx, campo_a_columna, "nivel_formacion") or extra.get("nivelFormacion"),
                coordinacion=_valor_texto(fila_valores, idx, campo_a_columna, "coordinacion") or extra.get("coordinacion"),
                codigoPrograma=_valor_texto(fila_valores, idx, campo_a_columna, "codigo_programa") or extra.get("codigoPrograma"),
                fechaInicioLectiva=_valor_fecha(fila_valores, idx, campo_a_columna, "fecha_inicio_lectiva") or extra.get("fechaInicioLectiva"),
                fechaFinLectiva=_valor_fecha(fila_valores, idx, campo_a_columna, "fecha_fin_lectiva") or extra.get("fechaFinLectiva"),
                fechaFinProductiva=_valor_fecha(fila_valores, idx, campo_a_columna, "fecha_fin_productiva") or extra.get("fechaFinProductiva"),
                faseActual=_valor_entero(fila_valores, idx, campo_a_columna, "fase_actual") or extra.get("faseActual"),
                faseActualEnBD=fase_actual_en_bd,
            )
        )
        if len(filas) >= MAX_FILAS_PREVIA:
            break

    advertencia_general = None
    if filas and all(f.idFicha is None for f in filas):
        if "ficha" not in campo_a_columna:
            advertencia_general = (
                "No se encontró ninguna columna de ficha reconocible en este archivo. Si es un "
                "formato de matriz o pivote (varias filas por ficha, ej. una fila de temas, otra de "
                "instructor y otra de ambiente por cada ficha), el asistente todavía no lo soporta -- "
                "usa un archivo con una fila por ficha."
            )
        else:
            advertencia_general = (
                "Se encontró una columna de ficha, pero ningún valor de esa columna se pudo leer "
                "como un número de ficha -- revisa si el archivo trae el número de ficha mezclado con "
                "otro texto en la misma celda."
            )

    return ImportarExcelPreviewResponse(
        nombreArchivo=nombre_archivo,
        hoja=ws.title,
        filaEncabezado=fila_encabezado,
        columnas=columnas,
        filas=filas,
        totalFilas=len(filas),
        filasConAdvertencia=sum(1 for f in filas if f.advertencia),
        advertenciaGeneral=advertencia_general,
        archivoComplementario=nombre_complementario if contenido_complementario else None,
        hojaComplementaria=hoja_complementaria,
    )


class _CatalogoPrograma:
    """Todo lo que `generar_propuesta` necesita de la BD, precargado en
    puñado de queries (una por tabla) en vez de una por ficha. Con un
    lote de 50 fichas, la versión anterior (una `_resultados_pendientes`
    por ficha, cada una con 3 queries propias + 2 más en el loop de
    `generar_propuesta`) hacía ~300 idas y vueltas a la BD -- contra
    Supabase remoto (no la BD en memoria de los tests) eso solo, sin
    tocar el solver, ya tardaba minutos. Encontrado en vivo el
    2026-09-13 probando con las 50 fichas reales: el batching por ficha
    del solver (ver _generar_bloques_por_ficha) no alcanzaba porque el
    cuello de botella real estaba ANTES de llegar al solver."""

    def __init__(self, db: Session, ids_ficha: list[int], id_trimestre: int):
        fichas = db.query(Ficha).filter(Ficha.idFicha.in_(ids_ficha)).all()
        self.fichas_por_id: dict[int, Ficha] = {f.idFicha: f for f in fichas}

        ids_programa = {f.idPrograma for f in fichas}
        competencias = (
            db.query(CompetenciaFormacion).filter(CompetenciaFormacion.idPrograma.in_(ids_programa)).all()
            if ids_programa
            else []
        )
        self.competencias_por_programa: dict[int, list[CompetenciaFormacion]] = {}
        for c in competencias:
            self.competencias_por_programa.setdefault(c.idPrograma, []).append(c)

        ids_competencia = [c.idCompetencia for c in competencias]
        resultados = (
            db.query(ResultadoAprendizaje).filter(ResultadoAprendizaje.idCompetencia.in_(ids_competencia)).all()
            if ids_competencia
            else []
        )
        self.resultados_por_competencia: dict[int, list[ResultadoAprendizaje]] = {}
        for r in resultados:
            self.resultados_por_competencia.setdefault(r.idCompetencia, []).append(r)

        horarios = (
            db.query(Horario)
            .filter(Horario.idFicha.in_(ids_ficha), Horario.idTrimestre == id_trimestre, Horario.activo.is_(True))
            .all()
            if ids_ficha
            else []
        )
        self.ids_resultado_con_horario_por_ficha: dict[int, set[int]] = {}
        for h in horarios:
            self.ids_resultado_con_horario_por_ficha.setdefault(h.idFicha, set()).add(h.idResultado)

        # Slots ya ocupados por horarios YA GUARDADOS este trimestre --
        # de CUALQUIER ficha, no solo las de este lote, porque un
        # instructor/ambiente ya reservado por una ficha ajena al lote
        # sigue sin estar libre. Encontrado en vivo el 2026-09-13: sin
        # esto, el solver proponía bloques que chocaban con horarios
        # reales ya guardados (revalidar con /horarios/validar los
        # rechazaba con "cruce_ficha"/"cruce_instructor"/"cruce_ambiente"
        # después de que el coordinador ya había visto la propuesta).
        # `generar_horario` espera (id_recurso, (horaInicio, horaFin),
        # idDia) -- ver _SlotOcupado en app/scheduling/generator.py.
        filas_dias = (
            db.query(
                Horario.idInstructor, Horario.idAmbiente, Horario.idFicha,
                Horario.horaInicio, Horario.horaFin, horario_dia.c.idDia,
            )
            .join(horario_dia, horario_dia.c.idHorario == Horario.idHorario)
            .filter(Horario.idTrimestre == id_trimestre, Horario.activo.is_(True))
            .all()
        )
        self.ocupados_instructor: set[tuple[str, tuple, int]] = set()
        self.ocupados_ambiente: set[tuple[int, tuple, int]] = set()
        self.ocupados_ficha: set[tuple[int, tuple, int]] = set()
        for id_instructor, id_ambiente, id_ficha, hora_inicio, hora_fin, id_dia in filas_dias:
            franja = (hora_inicio, hora_fin)
            self.ocupados_instructor.add((str(id_instructor), franja, id_dia))
            self.ocupados_ambiente.add((id_ambiente, franja, id_dia))
            self.ocupados_ficha.add((id_ficha, franja, id_dia))

        ambientes_disponibles = db.query(Ambiente).filter(Ambiente.estado_ambiente == "disponible").all()
        self.ambientes_por_id: dict[int, Ambiente] = {a.id: a for a in ambientes_disponibles}
        self.ids_ambiente_por_sede: dict[int, list[int]] = {}
        for a in ambientes_disponibles:
            self.ids_ambiente_por_sede.setdefault(a.sede_id, []).append(a.id)
        self.ids_ambiente_todos: list[int] = [a.id for a in ambientes_disponibles]

    def tiene_algun_resultado_definido(self, ficha: Ficha) -> bool:
        return bool(self.competencias_por_programa.get(ficha.idPrograma))

    def ids_ambiente_de(self, ficha: Ficha) -> list[int]:
        if ficha.idSede:
            return self.ids_ambiente_por_sede.get(ficha.idSede, [])
        return self.ids_ambiente_todos

    def resultados_pendientes(self, ficha: Ficha) -> list[ResultadoAprendizaje]:
        """Resultados de aprendizaje del programa de la ficha que todavía
        no tienen un horario activo en este trimestre -- son los que hay
        que programar.

        Si la ficha tiene `faseActual` (en qué fase de SU pénsum va,
        1=TRIM I..4=TRIM IV -- ver PLAN_INTEGRACION_IA.md), solo se traen
        los resultados de esa fase. Programas cuyo currículo no trae fase
        (`numeroFase` nulo en todos los resultados, import viejo desde la
        hoja combinada) siguen sin filtrarse -- si se filtrara iban a
        desaparecer todos por no calzar con ningún número.

        También se ignora `faseActual` si NO coincide con NINGUNA fase
        real del currículo cargado (ej. llega 5, 6 o 7 desde la columna
        "TRM" de un archivo -- confirmado que puede traer números romanos
        hasta VII -- pero el currículo de este programa solo tiene fases
        1-4 cargadas). Encontrado en vivo el 2026-09-13 cruzando dos
        archivos reales que traen la fase con criterios distintos (TRI de
        uno, TRM de otro, no siempre de acuerdo entre sí): sin esta
        guarda, una ficha con faseActual fuera de rango se marcaba en
        silencio como "ya tiene todo programado" (cero resultados
        pendientes) en vez de tratarse como dato no confiable."""
        competencias = self.competencias_por_programa.get(ficha.idPrograma, [])
        resultados = [
            r for c in competencias for r in self.resultados_por_competencia.get(c.idCompetencia, [])
        ]

        fases_del_curriculo = {r.numeroFase for r in resultados if r.numeroFase is not None}
        if ficha.faseActual is not None and ficha.faseActual in fases_del_curriculo:
            resultados = [r for r in resultados if r.numeroFase == ficha.faseActual]

        ids_con_horario = self.ids_resultado_con_horario_por_ficha.get(ficha.idFicha, set())
        return [r for r in resultados if r.idResultado not in ids_con_horario]


def _diagnostico_infactibilidad(necesidades: list[NecesidadHorario], jornada: str) -> str:
    """Mensaje determinista (sin IA -- las restricciones duras no las
    decide un LLM) cuando el solver no encuentra combinación sin choques.
    No repite exactamente por qué CP-SAT falló (eso es una prueba de
    infactibilidad, no algo que valga la pena mostrarle al coordinador);
    en cambio da una cota simple de capacidad vs demanda para que sepa
    qué palanca mover: menos fichas/resultados por lote, más
    instructores/ambientes, o definir la fase actual de cada ficha para
    no intentar programar currículo de trimestres que todavía no tocan."""
    dias_por_semana = 5
    franjas = len(FRANJAS_POR_JORNADA.get(jornada, []))

    ids_instructor: set[str] = set()
    ids_ambiente: set[int] = set()
    for n in necesidades:
        ids_instructor.update(n.instructores_candidatos)
        ids_ambiente.update(n.ambientes_candidatos)

    capacidad_instructor = len(ids_instructor) * franjas * dias_por_semana
    capacidad_ambiente = len(ids_ambiente) * franjas * dias_por_semana
    demanda = len(necesidades)

    partes = [
        f"No se encontró una combinación sin choques: hay {demanda} resultado(s) por programar en la "
        f"jornada {jornada}, pero la capacidad esa semana es de {capacidad_instructor} bloque(s) con los "
        f"{len(ids_instructor)} instructor(es) disponibles y {capacidad_ambiente} con los {len(ids_ambiente)} "
        "ambiente(s) disponibles (franjas × 5 días)."
    ]
    if demanda > min(capacidad_instructor, capacidad_ambiente):
        partes.append(
            "Reduce cuántas fichas o resultados generas juntos en un mismo lote, o define la fase "
            "actual del pénsum de cada ficha (en Fichas) si el programa ya tiene su currículo dividido "
            "por trimestre -- así solo se programan los resultados que tocan ahora, no todo el programa."
        )
    return " ".join(partes)


# Tope de instructores/ambientes candidatos que se le ofrecen al solver
# POR NECESIDAD. `generar_horario` arma una opción por cada combinación
# de franja × patrón de día × instructor × ambiente -- con catálogos
# reales (cientos de instructores/ambientes, no los 4-6 de prueba) ese
# producto cruzado explota a cientos de millones de variables antes de
# construir el modelo siquiera. No hace falta ofrecerle al solver los
# 215 instructores del centro para dictar una sola clase: un subconjunto
# rotado (distinto por necesidad, para no pelear todas por los mismos
# candidatos) alcanza de sobra. Ver PLAN_INTEGRACION_IA.md.
#
# Bajado de 8 a 6 el 2026-09-12 al medir el caso real de abajo: con 8,
# resolver ficha por ficha (ver _generar_bloques_por_ficha) tardaba ~44s
# para 50 fichas -- al límite del timeout del frontend (45s). Con 6 baja
# a ~13-15s con catálogos realistas (200+ instructores) sin perder
# factibilidad en las pruebas.
_LIMITE_CANDIDATOS = 6

# Tope de necesidades de UNA SOLA ficha. `generar_horario` sigue siendo
# O(necesidades × opciones) *dentro de un mismo lote* -- ver
# _generar_bloques_por_ficha, que ya resuelve ficha por ficha (no todas
# las fichas juntas en un solo modelo) para que esto casi nunca se
# dispare. Sigue existiendo como salvavidas para el caso patológico de
# UNA ficha sin faseActual cuyo programa trae cientos de resultados
# pendientes (currículo completo, no solo el trimestre que toca): esa
# ficha se salta y se reporta en `fichasSinProgramar` en vez de intentar
# construir un modelo gigante para ella sola.
_MAX_NECESIDADES_POR_FICHA = 300

# Presupuesto de tiempo total para generar la propuesta completa (todas
# las fichas del lote). Encontrado en vivo el 2026-09-12: ~50 fichas sin
# faseActual seleccionadas juntas dejaban el request colgado 8+ minutos
# hasta que el frontend hacía timeout (45s) sin explicar por qué.
# Resolver ficha por ficha (en vez de un solo modelo con todas las
# necesidades) evita eso en el caso normal, pero si el catálogo de
# instructores/ambientes es chico y muchas fichas tardan en resolverse
# igual, este tope corta el lote a tiempo: lo que ya se resolvió se
# devuelve como propuesta parcial (factible=True) y el resto queda en
# `fichasSinProgramar` para reintentar en un lote más chico, en vez de
# dejar que el frontend haga timeout sin ninguna propuesta.
_PRESUPUESTO_TIEMPO_TOTAL_SEG = 30.0
_TIEMPO_LIMITE_SOLVER_POR_FICHA_SEG = 5.0


def _muestra_rotada(candidatos: list, tamano: int, offset: int) -> list:
    if len(candidatos) <= tamano:
        return candidatos
    inicio = offset % len(candidatos)
    rotado = candidatos[inicio:] + candidatos[:inicio]
    return rotado[:tamano]


def generar_propuesta(
    db: Session, id_trimestre: int, ids_ficha: list[int], jornada: str
) -> GenerarPropuestaResponse:
    if jornada not in JORNADAS_VALIDAS:
        raise ValueError(f"Jornada '{jornada}' no reconocida -- debe ser una de {JORNADAS_VALIDAS}.")

    instructores = db.query(Usuario).filter(Usuario.tipoContrato.isnot(None), Usuario.estado == "activo").all()
    if not instructores:
        return GenerarPropuestaResponse(bloques=[], factible=False, mensaje="No hay instructores activos en el catálogo.")
    ids_instructor = [str(i.idUsuario) for i in instructores]
    nombres_instructor = {str(i.idUsuario): i.nombre for i in instructores}

    catalogo = _CatalogoPrograma(db, ids_ficha, id_trimestre)

    necesidades_por_ficha: dict[int, list[NecesidadHorario]] = {}
    fichas_por_id: dict[int, Ficha] = catalogo.fichas_por_id
    resultados_por_id: dict[int, ResultadoAprendizaje] = {}
    total_necesidades = 0
    # Para distinguir "ya está todo programado" (tiene resultados, todos
    # con horario) de "este programa no tiene resultados de aprendizaje
    # definidos todavía" (nada que programar porque no hay currículo
    # cargado) -- son causas muy distintas y el mensaje debe decir cuál es.
    hay_programas_sin_resultados = False

    for id_ficha in ids_ficha:
        ficha = fichas_por_id.get(id_ficha)
        if not ficha:
            continue

        ids_ambiente = catalogo.ids_ambiente_de(ficha)
        if not ids_ambiente:
            continue

        if not catalogo.tiene_algun_resultado_definido(ficha):
            hay_programas_sin_resultados = True

        for resultado in catalogo.resultados_pendientes(ficha):
            resultados_por_id[resultado.idResultado] = resultado
            necesidades_por_ficha.setdefault(id_ficha, []).append(
                NecesidadHorario(
                    id_ficha=id_ficha,
                    id_resultado=resultado.idResultado,
                    jornada=jornada,
                    instructores_candidatos=_muestra_rotada(ids_instructor, _LIMITE_CANDIDATOS, total_necesidades),
                    ambientes_candidatos=_muestra_rotada(ids_ambiente, _LIMITE_CANDIDATOS, total_necesidades),
                )
            )
            total_necesidades += 1

    if not necesidades_por_ficha:
        if hay_programas_sin_resultados:
            mensaje = (
                "El programa de una o más de las fichas seleccionadas no tiene resultados de "
                "aprendizaje (competencias) definidos todavía -- no hay nada que programar hasta que "
                "se cargue ese contenido curricular."
            )
        else:
            mensaje = "Las fichas seleccionadas ya tienen todos sus resultados programados."
        return GenerarPropuestaResponse(bloques=[], factible=True, mensaje=mensaje)

    asignaciones, fichas_sin_programar = _generar_bloques_por_ficha(
        necesidades_por_ficha,
        catalogo.ocupados_instructor,
        catalogo.ocupados_ambiente,
        catalogo.ocupados_ficha,
    )

    if not asignaciones:
        primera_ficha = next(iter(necesidades_por_ficha))
        return GenerarPropuestaResponse(
            bloques=[], factible=False,
            mensaje=_diagnostico_infactibilidad(necesidades_por_ficha[primera_ficha], jornada)
            if len(necesidades_por_ficha) == 1
            else (
                f"No se pudo programar ninguna de las {len(necesidades_por_ficha)} ficha(s) sin choques -- "
                "probablemente no hay suficientes instructores o ambientes disponibles para tantas fichas "
                "juntas. Reduce cuántas fichas generas a la vez, o define la fase actual del pénsum de "
                "cada ficha en Fichas para que solo traiga los resultados que tocan ahora."
            ),
        )

    ambientes_por_id = catalogo.ambientes_por_id
    id_jornada = _ID_JORNADA_POR_CLAVE[jornada]

    bloques = [
        BloquePropuesto(
            idFicha=b.id_ficha,
            fichaCodigo=fichas_por_id[b.id_ficha].codigoFicha,
            idResultado=b.id_resultado,
            resultadoDescripcion=resultados_por_id[b.id_resultado].descripcion,
            idInstructor=b.id_instructor,
            instructorNombre=nombres_instructor[b.id_instructor],
            idAmbiente=b.id_ambiente,
            ambienteNombre=ambientes_por_id[b.id_ambiente].nombre,
            idJornada=id_jornada,
            dias=list(b.dias),
            horaInicio=b.hora_inicio,
            horaFin=b.hora_fin,
        )
        for b in asignaciones
    ]

    codigos_fichas_sin_programar = [fichas_por_id[fid].codigoFicha for fid in fichas_sin_programar]
    if codigos_fichas_sin_programar:
        # No todas fallan por lo mismo: una ficha con más necesidades que
        # slots hay en una semana (franjas × 5 días) no cabe SIN IMPORTAR
        # cuántos instructores/ambientes sobren -- eso casi siempre es
        # porque a la ficha le falta `faseActual` y le llegó el pénsum
        # completo del programa en vez de solo el trimestre que le toca
        # (ver _CatalogoPrograma.resultados_pendientes). Encontrado en
        # vivo el 2026-09-13: el mensaje genérico ("no quedan
        # instructores/ambientes") llevaba al coordinador a buscar en el
        # lugar equivocado para estas fichas.
        capacidad_semanal = len(FRANJAS_POR_JORNADA[jornada]) * 5
        sin_fase = [
            fichas_por_id[fid].codigoFicha
            for fid in fichas_sin_programar
            if fichas_por_id[fid].faseActual is None and len(necesidades_por_ficha[fid]) > capacidad_semanal
        ]
        sin_recursos = [c for c in codigos_fichas_sin_programar if c not in sin_fase]

        partes = [f"{len(bloques)} bloques propuestos."]
        if sin_fase:
            partes.append(
                f"{len(sin_fase)} ficha(s) ({', '.join(sin_fase)}) traen más resultados pendientes que "
                f"slots hay en una semana de esta jornada ({capacidad_semanal}) porque no tienen la fase "
                "actual del pénsum definida -- les llega el programa completo en vez de solo este "
                "trimestre. Defínela en Fichas (o en el archivo complementario del import) para que solo "
                "traiga lo que toca ahora."
            )
        if sin_recursos:
            partes.append(
                f"{len(sin_recursos)} ficha(s) ({', '.join(sin_recursos)}) no se pudieron programar sin "
                "choques -- probablemente ya no quedan instructores o ambientes libres para ellas en esta "
                "jornada. Genera esas fichas por separado, en otra jornada, o con más instructores/ambientes "
                "disponibles."
            )
        mensaje = " ".join(partes)
    else:
        mensaje = f"{len(bloques)} bloques propuestos."
    return GenerarPropuestaResponse(
        bloques=bloques, factible=True, mensaje=mensaje, fichasSinProgramar=codigos_fichas_sin_programar
    )


def _generar_bloques_por_ficha(
    necesidades_por_ficha: dict[int, list[NecesidadHorario]],
    ocupados_instructor: set,
    ocupados_ambiente: set,
    ocupados_ficha: set,
) -> tuple[list, list[int]]:
    """Resuelve el CP-SAT ficha por ficha en vez de todas las necesidades
    juntas en un solo modelo. Encontrado en vivo el 2026-09-12: ~50
    fichas sin faseActual seleccionadas juntas son miles de necesidades,
    y aunque _muestra_rotada ya acota las opciones POR necesidad,
    construir un solo modelo con todas ellas es demasiado lento incluso
    en minutos. Resolver ficha por ficha mantiene cada modelo chico
    (una ficha típica trae decenas de resultados, no miles) y sigue
    evitando choques entre fichas porque los instructores/ambientes/
    fichas que ya se usaron se acarrean como `ocupados_*` de una ficha a
    la siguiente -- ver generar_horario en app/scheduling/generator.py.

    `ocupados_*` vienen sembrados por el llamador con lo que YA está
    guardado en la BD para este trimestre (ver _CatalogoPrograma) -- así
    el solver nunca propone un bloque que choque con un horario real ya
    existente, no solo con lo que se va asignando dentro de este mismo
    lote. Se mutan in-place (el llamador los sigue usando después).

    Devuelve (bloques_asignados, ids_ficha_sin_programar). Una ficha
    puede quedar sin programar por infactibilidad genuina (ya no caben
    sus resultados sin chocar con lo que otras fichas -- o la BD -- ya
    ocuparon) o porque se agotó el presupuesto de tiempo total del lote
    -- en cualquier caso se reporta para que el coordinador la reintente
    en vez de dejar todo el request sin respuesta."""
    inicio = time.perf_counter()
    bloques: list = []
    fichas_sin_programar: list[int] = []

    ids_ficha = list(necesidades_por_ficha)
    for indice, id_ficha in enumerate(ids_ficha):
        necesidades_ficha = necesidades_por_ficha[id_ficha]

        if len(necesidades_ficha) > _MAX_NECESIDADES_POR_FICHA:
            fichas_sin_programar.append(id_ficha)
            continue

        if time.perf_counter() - inicio > _PRESUPUESTO_TIEMPO_TOTAL_SEG:
            fichas_sin_programar.extend(ids_ficha[indice:])
            break

        asignacion_ficha = generar_horario(
            necesidades_ficha,
            tiempo_limite_seg=_TIEMPO_LIMITE_SOLVER_POR_FICHA_SEG,
            ocupados_instructor=ocupados_instructor,
            ocupados_ambiente=ocupados_ambiente,
            ocupados_ficha=ocupados_ficha,
        )
        if asignacion_ficha is None:
            fichas_sin_programar.append(id_ficha)
            continue

        bloques.extend(asignacion_ficha)
        for b in asignacion_ficha:
            for dia in b.dias:
                ocupados_instructor.add((b.id_instructor, (b.hora_inicio, b.hora_fin), dia))
                ocupados_ambiente.add((b.id_ambiente, (b.hora_inicio, b.hora_fin), dia))
                ocupados_ficha.add((b.id_ficha, (b.hora_inicio, b.hora_fin), dia))

    return bloques, fichas_sin_programar
